#!/usr/bin/env python3
"""
Snapshot every Option Rider layer from Redis into a client-facing Excel workbook.

One row per module, columns: layer, module, check rule, inputs, then
per-symbol input values / output values / status.

Usage (pipeline must be running):
    python export_client_excel.py
    python export_client_excel.py --symbols RELIANCE,TCS,INFY
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import os
from pathlib import Path
from typing import Any, Dict, List, Optional

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

load_dotenv()

import redis

from app.config import REDIS_URL, load_symbols

try:
    from zoneinfo import ZoneInfo

    IST = ZoneInfo("Asia/Kolkata")
except Exception:
    IST = dt.timezone(dt.timedelta(hours=5, minutes=30))

OUT_DIR = Path(__file__).resolve().parent / "client_output"


def _now_ist() -> dt.datetime:
    return dt.datetime.now(IST)


def _j(obj: Any) -> str:
    if obj is None or obj == "" or obj == {}:
        return ""
    if isinstance(obj, str):
        return obj
    return json.dumps(obj, ensure_ascii=True, separators=(", ", ": "), default=str)


def _load_json(r: redis.Redis, key: str) -> Optional[Any]:
    raw = r.get(key)
    if not raw:
        return None
    try:
        return json.loads(raw)
    except Exception:
        return {"_raw": str(raw)[:500]}


def _pick_stream(r: redis.Redis, stream: str, symbol: str, n: int = 80) -> Optional[dict]:
    try:
        rows = r.xrevrange(stream, count=n)
    except Exception:
        return None
    want = symbol.upper()
    for _mid, fields in rows:
        for k in ("symbol", "underlying", "key"):
            v = str(fields.get(k) or "").upper()
            if v == want or v.startswith(want):
                return dict(fields)
    if rows:
        # last message on stream (regime is market-wide)
        return dict(rows[0][1])
    return None


def _status(out_val: str) -> str:
    if not out_val or out_val in ("{}", "null", "[]"):
        return "EMPTY"
    return "OK"


def _age_s(doc: Any) -> str:
    if not isinstance(doc, dict):
        return ""
    ts = doc.get("ts_ms") or doc.get("bar_ts_ms")
    try:
        age = time_now_ms() / 1000.0 - int(ts) / 1000.0
        if age < 90:
            return f"{age:.0f}s ago"
        if age < 7200:
            return f"{age / 60:.1f}m ago"
        return f"{age / 3600:.1f}h ago"
    except Exception:
        return ""


def time_now_ms() -> int:
    return int(_now_ist().timestamp() * 1000)


def collect_symbol(r: redis.Redis, sym: str) -> Dict[str, Any]:
    s = sym.upper()
    tick = _pick_stream(r, "md:ticks:eq", s, 120)
    opt_tick = _pick_stream(r, "md:ticks:opt", s, 200)
    c1m = _pick_stream(r, "md:candles:1m", s, 40)
    c5m = _pick_stream(r, "md:candles:5m", s, 20)
    c10m = _pick_stream(r, "md:candles:10m", s, 20)
    c30m = _pick_stream(r, "md:candles:30m", s, 20)
    c1d = _pick_stream(r, "md:candles:1d", s, 20)
    vol_blob = _load_json(r, "md:volume:latest") or {}
    vol = vol_blob.get(s) if isinstance(vol_blob, dict) else None
    if not isinstance(vol, dict):
        vol = {}
    regime = _load_json(r, "md:regime:latest") or {}
    expiry = r.hgetall("md:active_expiry") or {}

    return {
        "tick": tick or {},
        "opt_tick": opt_tick or {},
        "c1m": c1m or {},
        "c5m": c5m or {},
        "c10m": c10m or {},
        "c30m": c30m or {},
        "c1d": c1d or {},
        "expiry": expiry.get(s, ""),
        "st": _load_json(r, f"md:supertrend:bias:latest:{s}") or {},
        "ema": _load_json(r, f"md:ema:cross:latest:{s}") or {},
        "htf": _load_json(r, f"md:htf:trend:latest:{s}") or {},
        "pivots": _load_json(r, f"md:pivots:prevday:{s}") or {},
        "level": _load_json(r, f"md:level:entry:latest:{s}") or {},
        "momentum": _load_json(r, f"md:momentum:confirm:latest:{s}") or {},
        "volume": vol,
        "regime": regime if isinstance(regime, dict) else {},
        "bidask": _load_json(r, f"md:bidask:latest:{s}") or {},
        "smartmoney": _load_json(r, f"md:smartmoney:latest:{s}") or {},
        "orderflow": _load_json(r, f"md:orderflow:latest:{s}") or {},
        "imbalance": _load_json(r, f"md:imbalance:latest:{s}") or {},
        "stockflow": _load_json(r, f"md:stockflow:latest:{s}") or {},
        "composite": _load_json(r, f"md:composite:latest:{s}") or {},
        "oi_und": _load_json(r, f"md:oi:underlying:latest:{s}") or {},
        "greeks_ce": _load_json(r, f"md:greeks:phase:underlying:latest:{s}:CE") or {},
        "greeks_pe": _load_json(r, f"md:greeks:phase:underlying:latest:{s}:PE") or {},
        "strikeflow": _load_json(r, f"md:strikeflow:latest:{s}") or {},
        "expected": _load_json(r, f"md:expected_move:latest:{s}") or {},
        "entry": _load_json(r, f"md:entry:trigger:latest:{s}") or {},
        "strike": _load_json(r, f"md:strike:select:latest:{s}") or {},
        "capital": _load_json(r, f"md:capital:alloc:latest:{s}") or {},
    }


def _liq_for_underlying(r: redis.Redis, sym: str) -> dict:
    """Pick ATM-ish liquidity score using strikeflow/oi atm tradingsymbol if present."""
    sf = _load_json(r, f"md:strikeflow:latest:{sym}") or {}
    for k in ("atm_ce", "atm_call", "call_tsym", "tradingsymbol"):
        tsym = str(sf.get(k) or "")
        if tsym:
            doc = _load_json(r, f"md:liquidity:score:latest:{tsym}")
            if doc:
                return doc
    gp = _load_json(r, f"md:greeks:phase:underlying:latest:{sym}:CE") or {}
    tsym = str(gp.get("tradingsymbol") or "")
    if tsym:
        doc = _load_json(r, f"md:liquidity:score:latest:{tsym}")
        if doc:
            return doc
    return {}


def _greeks_change(r: redis.Redis, strike_doc: dict) -> dict:
    tsym = str((strike_doc or {}).get("tradingsymbol") or "")
    if not tsym:
        return {}
    return _load_json(r, f"md:greeks_change:latest:{tsym}") or {}


def build_rows(r: redis.Redis, symbols: List[str], data: Dict[str, dict]) -> List[dict]:
    def trio_in(getter) -> Dict[str, str]:
        return {s: _j(getter(data[s])) for s in symbols}

    def trio_out(getter) -> Dict[str, str]:
        return {s: _j(getter(data[s])) for s in symbols}

    greeks_snap_n = 0
    try:
        greeks_snap_n = int(r.xlen("md:greeks:snap"))
    except Exception:
        pass

    modules = [
        (
            "0 Data Ingestion",
            "Producer — equity ticks (Angel WS)",
            "LTP should match NSE; bid < ask.",
            "Angel SNAP_QUOTE: last_traded_price, bid, ask, volume",
            lambda d: d["tick"],
            "symbol, ltp/last_traded_price, bid, ask, volume",
            lambda d: {
                k: d["tick"].get(k)
                for k in d["tick"]
                if k in (
                    "symbol", "ltp", "last_traded_price", "bid", "ask",
                    "best_bid", "best_ask", "volume", "volume_trade_for_the_day",
                    "ts_ms", "ts_recv",
                )
            } or d["tick"],
        ),
        (
            "0 Data Ingestion",
            "Producer — option ticks + active expiry",
            "Expiry should be the near NFO weekly/monthly. Option prints should exist.",
            "ScripMaster ATM±STRIKES_AROUND for this underlying",
            lambda d: {"active_expiry": d["expiry"], "sample_opt": d["opt_tick"]},
            "active_expiry + last option tick for this underlying",
            lambda d: {"active_expiry": d["expiry"], "opt": d["opt_tick"]},
        ),
        (
            "0 Data Ingestion",
            "Greeks poller (REST md:greeks:snap)",
            f"Stream length now = {greeks_snap_n}. Non-zero means REST greeks arrived.",
            "md:active_expiry hash",
            lambda d: {"active_expiry": d["expiry"]},
            "md:greeks:snap message count (market-wide)",
            lambda d: {"greeks_snap_xlen": greeks_snap_n},
        ),
        (
            "1.1 Core — candles",
            "1-minute OHLCV",
            "high >= max(open,close); low <= min(open,close).",
            "Equity ticks aggregated to 1m",
            lambda d: d["tick"],
            "o, h, l, c, v, ts_ms",
            lambda d: d["c1m"],
        ),
        (
            "1.1 Core — candles",
            "5m / 10m / 30m / 1d candles",
            "Higher TFs fill only after those bars close. 1d empty on first live day is expected.",
            "1m candles resampled",
            lambda d: d["c1m"],
            "last 5m, 10m, 30m, 1d bars",
            lambda d: {"5m": d["c5m"], "10m": d["c10m"], "30m": d["c30m"], "1d": d["c1d"]},
        ),
        (
            "1.1 Indicator Signals",
            "Supertrend MTF (ATR 7, multiplier 1)",
            "bias=CALL if >=3 of 4 TFs bullish; PUT if >=3 bearish; else NEUTRAL. na = not enough bars.",
            "OHLC 1m/5m/10m/30m",
            lambda d: {"c1m": d["c1m"], "c5m": d["c5m"], "c10m": d["c10m"], "c30m": d["c30m"]},
            "bias, bullish count, bearish count, st_1m/5m/10m/30m",
            lambda d: d["st"],
        ),
        (
            "1.1 Indicator Signals",
            "EMA Cross 9/26 (1m)",
            "state=bullish iff ema9 > ema26. EMPTY until ~26 closed 1m bars.",
            "1m closes",
            lambda d: d["c1m"],
            "state, ema9, ema26, signal (cross or none)",
            lambda d: d["ema"],
        ),
        (
            "1.1 Indicator Signals",
            "Higher-timeframe filter (D/W/M)",
            "CALL only if daily AND weekly AND monthly close > previous period. Needs prior daily candles.",
            "md:candles:1d history",
            lambda d: d["c1d"],
            "bias, daily/weekly/monthly, close vs prev",
            lambda d: d["htf"],
        ),
        (
            "1.1 Indicator Signals",
            "Classic pivots (previous day H/L/C)",
            "P=(H+L+C)/3; R1=2P-L; S1=2P-H. EMPTY until a completed prior daily candle exists.",
            "Previous session high, low, close",
            lambda d: d["c1d"],
            "date, P, R1, R2, S1, S2",
            lambda d: d["pivots"],
        ),
        (
            "1.1 Indicator Signals",
            "Level entry (pivot / R1 / S1 break)",
            "BUY CALL if 1m close crosses up P or R1; BUY PUT crosses down P or S1. Needs pivots.",
            "1m close + prev-day pivots",
            lambda d: {"c1m": d["c1m"], "pivots": d["pivots"]},
            "signal, level, strength, reason",
            lambda d: d["level"],
        ),
        (
            "1.1 Indicator Signals",
            "Momentum confirm (Supertrend AND EMA)",
            "BUY CALL if ST CALL and EMA bullish; BUY PUT if ST PUT and EMA bearish.",
            "supertrend bias + ema state",
            lambda d: {"st": d["st"], "ema": d["ema"]},
            "confirmed side",
            lambda d: d["momentum"],
        ),
        (
            "1.1 Volume",
            "Volume analyzer (buyer vs seller on 1m range)",
            "buy% ≈ (close-low)/(high-low)*100. Bullish if buy%>=60; Bearish if sell%>=60; Strong if surge>2.",
            "1m high, low, close, volume",
            lambda d: {k: d["volume"].get(k) for k in ("high", "low", "close", "volume")} or d["c1m"],
            "buy_pct, sell_pct, volume_surge, signal",
            lambda d: d["volume"],
        ),
        (
            "1.1 Regime",
            "Market regime (advance/decline)",
            "With only 3 symbols this is NOT NSE breadth — only those names. Bullish if advance%>=60 or ratio>=1.5.",
            "All symbols in symbols.txt vs previous close",
            lambda d: {"universe": "symbols.txt"},
            "regime, advance_pct, breadth_ratio, capital bias",
            lambda d: d["regime"],
        ),
        (
            "1.1 Bid-Ask",
            "Spread / liquidity",
            "Liquid cash names: spread% typically << 0.05%. bid must be < ask.",
            "Top of book on equity (and options)",
            lambda d: d["tick"],
            "bid, ask, spread, signal",
            lambda d: d["bidask"],
        ),
        (
            "1.1 Bid-Ask",
            "Smart money (walls / absorption / sweep)",
            "Flags only — not a trade by itself.",
            "Tick sizes + book",
            lambda d: d["tick"],
            "wall / absorption / sweep fields",
            lambda d: d["smartmoney"],
        ),
        (
            "1.1 Bid-Ask",
            "Order flow (net delta)",
            "Positive rising net delta = buy pressure.",
            "Trades hitting bid vs ask",
            lambda d: d["tick"],
            "buy_pressure, sell_pressure, bias",
            lambda d: d["orderflow"],
        ),
        (
            "1.1 Bid-Ask",
            "Bid-ask quantity imbalance",
            "Imbalance > +0.30 bullish, < -0.30 bearish. Range -1 to +1.",
            "Bid qty vs ask qty (depth-weighted)",
            lambda d: d["tick"],
            "raw, weighted_filtered, final_score",
            lambda d: d["imbalance"],
        ),
        (
            "1.1 Bid-Ask",
            "Stock entry/exit gate (5 conditions)",
            "Entry needs all 5 stock bid-ask conditions; any exit condition flips.",
            "spread, imbalance, net delta, walls, last print",
            lambda d: {"bidask": d["bidask"], "imbalance": d["imbalance"], "orderflow": d["orderflow"]},
            "entry/exit status",
            lambda d: d["stockflow"],
        ),
        (
            "1.1 Bid-Ask",
            "Composite bid-ask score",
            "Score > +0.60 → Entry; < -0.40 → Exit. Liquidity-exit override always wins.",
            "imbalance, net delta, smart money, spread, S/R, option flow",
            lambda d: {
                "imbalance": d["imbalance"],
                "orderflow": d["orderflow"],
                "smartmoney": d["smartmoney"],
                "bidask": d["bidask"],
                "strikeflow": d["strikeflow"],
            },
            "score, status, components, override",
            lambda d: d["composite"],
        ),
        (
            "1.2 OI",
            "Open interest — underlying",
            "spot ≈ last LTP. ATM near spot. positioning = Long/Short buildup etc.",
            "Option chain OI + spot",
            lambda d: {"spot_tick": d["tick"], "expiry": d["expiry"]},
            "spot, atm, max_pain, support, resistance, positioning",
            lambda d: d["oi_und"],
        ),
        (
            "1.2 Greeks",
            "Greeks phase (Accumulation / Markup / Distribution)",
            "Markup = entry zone; Distribution = exit; Accumulation = no trade. ATM CE and PE.",
            "delta, gamma, theta, vega, IV, price vs prior resistance",
            lambda d: {"expiry": d["expiry"]},
            "ATM CE phase + ATM PE phase",
            lambda d: {"CE": d["greeks_ce"], "PE": d["greeks_pe"]},
        ),
        (
            "1.2 Liquidity",
            "Option liquidity score (0–100)",
            "75–100 green full size; 0–24 do not enter. Shown for ATM CE if available.",
            "OI, volume, spread, book depth",
            lambda d: {"oi": d["oi_und"], "bidask": d["bidask"]},
            "score, band, final_entry_size",
            lambda d: d.get("liquidity") or {},
        ),
        (
            "1.2 Strike flow",
            "Option order-flow strike scan",
            "Unusual Vol/OI and sweeps inform strike preference.",
            "Order-flow bias + option prints",
            lambda d: d["orderflow"],
            "atm, status, bias",
            lambda d: d["strikeflow"],
        ),
        (
            "1.3 Volatility",
            "Expected move (prediction only)",
            "Does not pick strike. Empty IV is a data gap until REST greeks populate. indicator_score is Supertrend + EMA + pivot strength (-2..+2).",
            "spot, volume score, OI score, imbalance, IV if any",
            lambda d: {
                "spot": (d["tick"] or {}).get("ltp") or (d["oi_und"] or {}).get("spot"),
                "volume": d["volume"],
                "oi": d["oi_und"],
                "imbalance": d["imbalance"],
            },
            "expected move, range, direction, confidence/flags",
            lambda d: d["expected"],
        ),
        (
            "1.4 Decision",
            "Entry trigger (final CALL/PUT gate)",
            "BUY only if HTF + Supertrend + EMA + pivot break + volume + OI positioning + Greeks MARKUP all agree. reason explains NEUTRAL.",
            "htf, st, ema, level, volume, oi, greeks phase",
            lambda d: {
                "htf": d["htf"],
                "st": d["st"],
                "ema": d["ema"],
                "level": d["level"],
                "volume": d["volume"],
                "oi": d["oi_und"],
                "greeks_ce": d["greeks_ce"],
                "greeks_pe": d["greeks_pe"],
            },
            "signal, strength, level, reason, aligned",
            lambda d: d["entry"],
        ),
        (
            "1.4 Decision",
            "Strike selector",
            "Runs only after a BUY CALL / BUY PUT. Picks ATM or slight OTM (or OI target strike).",
            "entry_trigger BUY + option chain + spot",
            lambda d: d["entry"],
            "status, tradingsymbol, strike, atm, expiry, reason",
            lambda d: d["strike"],
        ),
        (
            "1.3 Volatility",
            "Greeks change predictor",
            "Needs a selected option (strike_select OK) plus expected move + IV. Empty until a BUY fires.",
            "strike_select + expected_move + greeks phase",
            lambda d: {"strike": d["strike"], "expected": d["expected"]},
            "greeks_change payload for selected tradingsymbol",
            lambda d: d.get("greeks_change") or {},
        ),
        (
            "1.6 Capital",
            "Capital allocation",
            "Sizes notional from regime 70/30 CALL/PUT split and 1–5% risk cap. Empty until strike_select OK.",
            "strike_select + md:regime:latest",
            lambda d: {"strike": d["strike"], "regime": d["regime"]},
            "side, budget, trade_notional, reason",
            lambda d: d["capital"],
        ),
        (
            "Backlog",
            "Lot sizing / ranking / order executor / kill switch / risk / SL / slippage / journal",
            "Named in Option-rider-algo-specs.md but not implemented — no live output.",
            "n/a",
            lambda d: {},
            "n/a",
            lambda d: {"status": "NOT IMPLEMENTED"},
        ),
    ]

    rows: List[dict] = []
    for layer, module, check, in_desc, in_fn, out_desc, out_fn in modules:
        row = {
            "Layer": layer,
            "Module": module,
            "How the client should verify": check,
            "What goes IN (description)": in_desc,
            "What comes OUT (description)": out_desc,
        }
        for s in symbols:
            inn = in_fn(data[s])
            out = out_fn(data[s])
            row[f"{s} INPUT"] = _j(inn)
            row[f"{s} OUTPUT"] = _j(out)
            st = "NOT IMPLEMENTED" if "NOT IMPLEMENTED" in _j(out) else _status(_j(out))
            if st == "OK":
                age = _age_s(out if isinstance(out, dict) else {})
                if age:
                    st = f"OK ({age})"
            row[f"{s} STATUS"] = st
        rows.append(row)
    return rows


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
TITLE_FONT = Font(bold=True, name="Calibri", size=16, color="1F4E79")
SECTION_FONT = Font(bold=True, name="Calibri", size=12, color="1F4E79")
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
OK_FILL = PatternFill("solid", fgColor="C6EFCE")
EMPTY_FILL = PatternFill("solid", fgColor="FCE4D6")
NA_FILL = PatternFill("solid", fgColor="D9D9D9")
ALT_FILL = PatternFill("solid", fgColor="F2F2F2")


def _style_header(ws: Worksheet, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _autosize(ws: Worksheet, widths: Dict[int, int]) -> None:
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32


def write_cover(wb: Workbook, symbols: List[str], when: dt.datetime, path: Path) -> None:
    ws = wb.active
    ws.title = "00_Cover"
    lines = [
        ("Option Rider — Layer audit for client review", TITLE_FONT),
        ("", None),
        (f"Snapshot (Asia/Kolkata): {when.strftime('%Y-%m-%d %H:%M:%S %Z')}", None),
        (f"Symbols: {', '.join(symbols)}", None),
        (f"File: {path.name}", None),
        ("", None),
        ("How to read this workbook", SECTION_FONT),
        (
            "Each data sheet lists every live algo layer. For each layer you get: "
            "what data goes IN, what the module writes OUT, then one INPUT / OUTPUT / STATUS "
            "column per symbol. STATUS=OK means Redis had a payload at snapshot time. "
            "STATUS=EMPTY means the layer has not produced a value yet (warmup, missing previous-day candle, or no BUY signal).",
            None,
        ),
        ("", None),
        ("What EMPTY usually means on a first live session", SECTION_FONT),
        ("• EMA 9/26 needs ~26 closed 1-minute bars.", None),
        ("• 5m / 10m / 30m Supertrend needs those candles to close.", None),
        ("• Previous-day pivots and D/W/M trend need a completed daily bar (often tomorrow).", None),
        ("• Entry trigger BUY CALL/PUT needs ALL of: HTF + Supertrend + EMA + pivot break + volume + OI + Greeks MARKUP. NEUTRAL with a reason is a valid output.", None),
        ("• Strike / capital / greeks-change stay empty until a BUY fires.", None),
        ("• Market regime on 3 names is NOT full NSE breadth — only those three stocks.", None),
        ("", None),
        ("How to verify (spot checks)", SECTION_FONT),
        ("1. Equity LTP vs NSE/TradingView for the same second.", None),
        ("2. 1m candle: high ≥ close ≥ low.", None),
        ("3. Volume buy% ≈ (close − low) / (high − low) × 100.", None),
        ("4. Supertrend: CALL only if at least 3 of 1m/5m/10m/30m are bullish (ATR 7, multiplier 1).", None),
        ("5. EMA bullish only if EMA9 > EMA26 on 1m close.", None),
        ("6. Pivot P = (prev high + prev low + prev close) / 3.", None),
        ("7. Bid < Ask on liquid names; cash spread% is typically well under 0.05%.", None),
        ("8. OI spot ≈ LTP; ATM strike near spot.", None),
        ("", None),
        ("Sheets", SECTION_FONT),
        ("01_Summary — one line per module, key output + status per symbol.", None),
        ("02_Layer_IO — full input JSON and output JSON per symbol (audit sheet).", None),
        ("03_Per_Symbol — one block per symbol, latest values flattened.", None),
        ("", None),
        ("Not implemented (spec backlog — no live data): lot sizing, probability engine, trade ranking, order executor, kill switch, risk, smart SL, slippage, portfolio monitor, trade journal.", None),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        cell = ws.cell(i, 1, text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if font:
            cell.font = font
        else:
            cell.font = Font(name="Calibri", size=11)
        ws.row_dimensions[i].height = 18 if len(text) < 80 else 48
    ws.column_dimensions["A"].width = 120


def write_summary(wb: Workbook, symbols: List[str], rows: List[dict]) -> None:
    ws = wb.create_sheet("01_Summary")
    headers = ["Layer", "Module", "How the client should verify"]
    for s in symbols:
        headers += [f"{s} key output", f"{s} status"]
    ws.append(headers)
    _style_header(ws, len(headers))
    for i, row in enumerate(rows, start=2):
        vals = [row["Layer"], row["Module"], row["How the client should verify"]]
        for s in symbols:
            out = row.get(f"{s} OUTPUT") or ""
            short = out if len(out) <= 400 else out[:397] + "..."
            vals += [short, row.get(f"{s} STATUS") or ""]
        ws.append(vals)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(i, c)
            cell.alignment = WRAP
            cell.border = THIN
            if i % 2 == 0 and c <= 3:
                cell.fill = ALT_FILL
        for j, s in enumerate(symbols):
            st_col = 4 + j * 2 + 1
            st = (row.get(f"{s} STATUS") or "")
            fill = OK_FILL if st.startswith("OK") else (NA_FILL if "NOT" in st else EMPTY_FILL)
            ws.cell(i, st_col).fill = fill
        ws.row_dimensions[i].height = 60
    widths = {1: 22, 2: 42, 3: 55}
    col = 4
    for _s in symbols:
        widths[col] = 48
        widths[col + 1] = 16
        col += 2
    _autosize(ws, widths)
    ws.sheet_view.showGridLines = False


def write_io(wb: Workbook, symbols: List[str], rows: List[dict]) -> None:
    ws = wb.create_sheet("02_Layer_IO")
    headers = [
        "Layer",
        "Module",
        "How the client should verify",
        "What goes IN (description)",
        "What comes OUT (description)",
    ]
    for s in symbols:
        headers += [f"{s} INPUT", f"{s} OUTPUT", f"{s} STATUS"]
    ws.append(headers)
    _style_header(ws, len(headers))
    for i, row in enumerate(rows, start=2):
        vals = [
            row["Layer"],
            row["Module"],
            row["How the client should verify"],
            row["What goes IN (description)"],
            row["What comes OUT (description)"],
        ]
        for s in symbols:
            vals += [row.get(f"{s} INPUT") or "", row.get(f"{s} OUTPUT") or "", row.get(f"{s} STATUS") or ""]
        ws.append(vals)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(i, c)
            cell.alignment = WRAP
            cell.border = THIN
        for j, s in enumerate(symbols):
            st_col = 6 + j * 3 + 2
            st = row.get(f"{s} STATUS") or ""
            fill = OK_FILL if st.startswith("OK") else (NA_FILL if "NOT" in st else EMPTY_FILL)
            ws.cell(i, st_col).fill = fill
        ws.row_dimensions[i].height = 90
    widths = {1: 20, 2: 36, 3: 42, 4: 36, 5: 36}
    col = 6
    for _s in symbols:
        widths[col] = 42
        widths[col + 1] = 42
        widths[col + 2] = 16
        col += 3
    _autosize(ws, widths)


def write_per_symbol(wb: Workbook, symbols: List[str], data: Dict[str, dict]) -> None:
    ws = wb.create_sheet("03_Per_Symbol")
    headers = ["Symbol", "Field group", "Field", "Value", "Age / notes"]
    ws.append(headers)
    _style_header(ws, 5)
    groups = [
        ("Last equity tick", "tick"),
        ("Last 1m candle", "c1m"),
        ("Supertrend MTF", "st"),
        ("EMA 9/26", "ema"),
        ("HTF D/W/M", "htf"),
        ("Prev-day pivots", "pivots"),
        ("Level entry", "level"),
        ("Momentum confirm", "momentum"),
        ("Volume", "volume"),
        ("Bid-ask", "bidask"),
        ("Imbalance", "imbalance"),
        ("Composite score", "composite"),
        ("OI underlying", "oi_und"),
        ("Greeks phase CE", "greeks_ce"),
        ("Greeks phase PE", "greeks_pe"),
        ("Expected move", "expected"),
        ("Entry trigger", "entry"),
        ("Strike select", "strike"),
        ("Capital alloc", "capital"),
        ("Liquidity (ATM CE if found)", "liquidity"),
        ("Greeks change", "greeks_change"),
        ("Regime (shared)", "regime"),
        ("Active expiry", "expiry"),
    ]
    r = 2
    for s in symbols:
        d = data[s]
        first = True
        for gname, key in groups:
            val = d.get(key)
            if isinstance(val, dict) and val:
                for fk, fv in val.items():
                    ws.cell(r, 1, s if first else "")
                    ws.cell(r, 2, gname)
                    ws.cell(r, 3, str(fk))
                    ws.cell(r, 4, "" if fv is None else str(fv))
                    ws.cell(r, 5, _age_s(val) if first else "")
                    first = False
                    for c in range(1, 6):
                        ws.cell(r, c).alignment = WRAP
                        ws.cell(r, c).border = THIN
                    r += 1
            else:
                ws.cell(r, 1, s if first else "")
                ws.cell(r, 2, gname)
                ws.cell(r, 3, "(no payload)")
                ws.cell(r, 4, "" if val in (None, {}, "") else str(val))
                ws.cell(r, 5, "EMPTY")
                ws.cell(r, 5).fill = EMPTY_FILL
                first = False
                for c in range(1, 6):
                    ws.cell(r, c).alignment = WRAP
                    ws.cell(r, c).border = THIN
                r += 1
        r += 1
    _autosize(ws, {1: 14, 2: 32, 3: 28, 4: 70, 5: 16})


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--symbols", default="", help="Comma list; default symbols.txt")
    args = parser.parse_args()

    if args.symbols.strip():
        symbols = [x.strip().upper() for x in args.symbols.split(",") if x.strip()]
    else:
        symbols = [s.upper() for s in load_symbols()]

    r = redis.Redis.from_url(REDIS_URL, decode_responses=True)
    r.ping()

    data: Dict[str, dict] = {}
    for s in symbols:
        data[s] = collect_symbol(r, s)
        data[s]["liquidity"] = _liq_for_underlying(r, s)
        data[s]["greeks_change"] = _greeks_change(r, data[s].get("strike") or {})

    when = _now_ist()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / f"OptionRider_LayerAudit_{when.strftime('%Y-%m-%d_%H%M')}.xlsx"

    rows = build_rows(r, symbols, data)
    wb = Workbook()
    write_cover(wb, symbols, when, path)
    write_summary(wb, symbols, rows)
    write_io(wb, symbols, rows)
    write_per_symbol(wb, symbols, data)
    wb.save(path)
    print(f"Wrote {path}")
    ok = sum(1 for row in rows if any(str(row.get(f"{s} STATUS", "")).startswith("OK") for s in symbols))
    empty = sum(1 for row in rows if all(not str(row.get(f"{s} STATUS", "")).startswith("OK") for s in symbols))
    print(f"Modules with at least one OK symbol: {ok}; all-empty or N/A: {empty}")


if __name__ == "__main__":
    main()
